from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document

from .models import (
    ExtractionWarning,
    NormalizedDocument,
    NormalizedSection,
    SourceContentType,
)

HEADING_STYLE_PATTERN = re.compile(r"^Heading\s+(\d+)$")
PARSER_VERSION = "parser-v1"


@dataclass
class _SectionAccumulator:
    heading: str
    heading_path: list[str]
    paragraphs: list[str] = field(default_factory=list)


def parse_docx(
    source_path: Path,
    *,
    source_uri: str,
    source_content_sha256: str,
) -> NormalizedDocument:
    document = Document(source_path)

    title = ""
    heading_stack: list[str] = []
    sections: list[NormalizedSection] = []
    warnings: list[ExtractionWarning] = [
        ExtractionWarning(
            code="DOCX_PAGE_LOCATIONS_UNAVAILABLE",
            message=(
                "DOCX parsing preserves heading structure, but page locations "
                "are not available from python-docx."
            ),
        )
    ]
    current_section: _SectionAccumulator | None = None

    def flush_current_section() -> None:
        nonlocal current_section

        if current_section is None or not current_section.paragraphs:
            return

        sections.append(
            NormalizedSection(
                section_id=f"section-{len(sections) + 1:03d}",
                heading=current_section.heading,
                heading_path=current_section.heading_path,
                text="\n\n".join(current_section.paragraphs),
            )
        )
        current_section = None

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if not text:
            continue

        style_name = paragraph.style.name if paragraph.style else ""
        heading_match = HEADING_STYLE_PATTERN.match(style_name)

        if heading_match:
            heading_level = int(heading_match.group(1))

            if heading_level == 1:
                flush_current_section()
                title = text
                heading_stack = [text]
                continue

            flush_current_section()
            heading_stack = heading_stack[: heading_level - 1]
            heading_stack.append(text)
            current_section = _SectionAccumulator(
                heading=text,
                heading_path=heading_stack.copy(),
            )
            continue

        if current_section is None:
            current_section = _SectionAccumulator(
                heading="Introduction",
                heading_path=["Introduction"],
            )

        current_section.paragraphs.append(text)

    flush_current_section()

    if not sections:
        raise ValueError("DOCX contained no extractable section text.")

    if document.tables:
        warnings.append(
            ExtractionWarning(
                code="DOCX_TABLES_NOT_EXTRACTED",
                message=(
                    "The document contains tables. Table extraction is not yet "
                    "implemented, so their content was not indexed."
                ),
            )
        )

    return NormalizedDocument(
        source_uri=source_uri,
        source_content_type=SourceContentType.DOCX,
        source_content_sha256=source_content_sha256,
        title=title or source_path.stem.replace("-", " ").title(),
        sections=sections,
        extraction_warnings=warnings,
        parser_version=PARSER_VERSION,
    )